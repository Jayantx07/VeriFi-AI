"""Fill the ABC (Aravalli Build Challenge, GDG Udaipur) submission template with VeriFi AI content.

Usage:  python submission/source/build_abc_deck.py
Input:  submission/ABC Submission Template.pptx
Output: submission/VeriFi-AI_ABC_Submission.pptx
Icons:  rendered by make_icons.js (react-icons: Simple Icons brand marks + Google Material icons)
"""
import copy
import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SRC = ROOT / "ABC Submission Template.pptx"
import os
OUT = Path(os.environ.get("DECK_OUT", ROOT / "VeriFi-AI_ABC_Submission.pptx"))
ICONS = HERE / "icons"

FONT = "DM Sans"
# Google palette: fills use the brand hues, text uses the darker accessible variants
BLUE, RED, YELLOW, GREEN = "4285F4", "EA4335", "FBBC04", "34A853"
TEXT_OF = {BLUE: "1A73E8", RED: "C5221F", YELLOW: "B06000", GREEN: "188038"}
TINT = {BLUE: "E8F0FE", RED: "FCE8E6", YELLOW: "FEF7E0", GREEN: "E6F4EA"}
INK, MUTED, SOFT, LINE, WHITE, GREY_BG = "202124", "5F6368", "80868B", "DADCE0", "FFFFFF", "F1F3F4"
ICON_INK = "3C4043"

ICON_NAMES = [
    # Google product marks (Simple Icons)
    "SiGooglecloud", "SiGooglegemini", "SiFirebase", "SiGmail", "SiGooglemeet", "SiGooglepubsub",
    "SiGooglecloudspanner", "SiGooglecloudstorage", "SiGooglebigquery", "SiFastapi", "SiNextdotjs",
    "SiWebauthn", "SiFidoalliance",
    # Google Material icons
    "MdEmail", "MdVideocam", "MdAccountBalance", "MdCurrencyRupee", "MdPerson", "MdVerifiedUser",
    "MdManageSearch", "MdAccountTree", "MdRule", "MdFingerprint", "MdGavel", "MdShield", "MdKey",
    "MdStorage", "MdPhoneIphone", "MdDashboard", "MdDescription", "MdWeb", "MdDns", "MdCloudUpload",
    "MdTune", "MdSwapHoriz", "MdPsychology", "MdGroups", "MdRocketLaunch", "MdLocationOn", "MdMap",
    "MdPublic", "MdVisibility", "MdTaskAlt", "MdChevronRight", "MdBuild", "MdTraffic", "MdArrowForward",
    "MdWarningAmber",
]


def ensure_icons():
    subprocess.run(["node", str(HERE / "make_icons.js"), str(ICONS), ",".join(ICON_NAMES)],
                   cwd=HERE, check=True)


def icon_path(name, color):
    return str(ICONS / f"{name}_{color}.png")


# ---------- primitives ----------
def shape(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def remove(slide, *sids):
    for sid in sids:
        el = shape(slide, sid)._element
        el.getparent().remove(el)


def place(shp, x=None, y=None, w=None, h=None):
    if x is not None: shp.left = Inches(x)
    if y is not None: shp.top = Inches(y)
    if w is not None: shp.width = Inches(w)
    if h is not None: shp.height = Inches(h)


def _set_color(rPr, hex_):
    fill = rPr.find(qn("a:solidFill"))
    if fill is None:
        fill = rPr.makeelement(qn("a:solidFill"), {})
        rPr.insert(0, fill)
    for c in list(fill):
        fill.remove(c)
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": hex_}))


def fill_text(shp, paras, para_tmpl=0):
    """Replace a template text box's text, keeping its paragraph and run formatting."""
    body = shp.text_frame._txBody
    ps = body.findall(qn("a:p"))
    tmpl_p = ps[para_tmpl]
    tmpl_r = tmpl_p.findall(qn("a:r"))[0]
    for p in ps:
        body.remove(p)
    for para in paras:
        p = copy.deepcopy(tmpl_p)
        for child in p.findall(qn("a:r")) + p.findall(qn("a:br")):
            p.remove(child)
        end = p.find(qn("a:endParaRPr"))
        runs = [(para, {})] if isinstance(para, str) else para
        for t, opts in runs:
            r = copy.deepcopy(tmpl_r)
            r.find(qn("a:t")).text = t
            rPr = r.find(qn("a:rPr"))
            if "b" in opts: rPr.set("b", "1" if opts["b"] else "0")
            if "sz" in opts: rPr.set("sz", str(int(opts["sz"] * 100)))
            if "color" in opts: _set_color(rPr, opts["color"])
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        body.append(p)


def rich(t, emph=INK):
    """Split 'plain **key phrase** plain' into runs; key phrases are bold and darker (the skim layer)."""
    return [(part, {"b": True, "color": emph} if i % 2 else {}) for i, part in enumerate(t.split("**")) if part]


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, line=1.0, italic=False, spc=0, name=None, emph=INK):
    """Add a text box. paras: list of str or list of (text, opts) runs; opts: b, sz, color, i.
    A str paragraph may mark key phrases with **...**; they render bold in `emph`."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if space_after: p.space_after = Pt(space_after)
        runs = rich(para, emph) if isinstance(para, str) else para
        for t, opts in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("sz", size))
            f.bold = opts.get("b", bold)
            f.italic = opts.get("i", italic)
            f.color.rgb = RGBColor.from_string(opts.get("color", color))
            if spc: r._r.get_or_add_rPr().set("spc", str(int(spc * 100)))
    return tb


def label(slide, x, y, t, color=TEXT_OF[BLUE], size=14, w=8.0):
    """Small uppercase section label with letter spacing."""
    return text(slide, x, y, w, 0.32, [t.upper()], size=size, color=color, bold=True, spc=1.5)


def shadow(shp, blur=14, dist=3, alpha=12):
    spPr = shp._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(Pt(blur)), dist=str(Pt(dist)),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, lift=False, dash=False, line_w=1.0,
        name=None, geom=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h))
    if geom == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(line_w)
        if dash: s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    if lift:
        shadow(s)
    else:
        s.shadow.inherit = False
    if name: s.name = name
    return s


def icon(slide, name, color, x, y, size):
    return slide.shapes.add_picture(icon_path(name, color), Inches(x), Inches(y), Inches(size), Inches(size))


def badge(slide, x, y, d, fill, icon_name, glyph=0.56, glyph_color=WHITE):
    """Material-style icon: white glyph in a Google-coloured circle (the deck's visual motif)."""
    box(slide, x, y, d, d, fill=fill, geom=MSO_SHAPE.OVAL, name=f"Badge {icon_name}")
    g = d * glyph
    icon(slide, icon_name, glyph_color, x + (d - g) / 2, y + (d - g) / 2, g)


def tw(t, size, bold=False):
    """Rough DM Sans text width in inches."""
    return len(t) * size / 72 * (0.56 if bold else 0.5)


def chip(slide, x, y, icon_name, icon_color, t, h=0.44, size=13, fill=WHITE, line=LINE, color=INK):
    w = 0.2 + h * 0.6 + 0.1 + tw(t, size, True) + 0.2
    box(slide, x, y, w, h, fill=fill, line=line, radius=0.5, name=f"Chip {t}")
    icon(slide, icon_name, icon_color, x + 0.18, y + h * 0.2, h * 0.6)
    tx = x + 0.18 + h * 0.6 + 0.1
    text(slide, tx, y, w - (tx - x), h, [t], size=size, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)
    return w


def arrow(slide, pts, color=SOFT, width=1.75, dash=None):
    """Polyline through pts (inches); arrowhead on the last segment."""
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(pts, pts[1:])):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = RGBColor.from_string(color)
        c.line.width = Pt(width)
        if dash: c.line.dash_style = dash
        if i == len(pts) - 2:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def google_badge(slide):
    """'Built on Google Cloud' pill, top-right of every content slide."""
    t = "Built on Google Cloud"
    w = 0.2 + 0.32 + 0.12 + tw(t, 13, True) + 0.22
    x = 19.06 - w
    box(slide, x, 0.5, w, 0.5, fill=WHITE, line=LINE, radius=0.5, lift=True, name="Google badge")
    icon(slide, "SiGooglecloud", BLUE, x + 0.2, 0.59, 0.32)
    text(slide, x + 0.64, 0.5, w - 0.64, 0.5, [t], size=13, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)


ensure_icons()
prs = Presentation(SRC)
s1, s2, s3, s4, s5, s6 = prs.slides

# =====================================================================
# 1. Title
# =====================================================================
fill_text(shape(s1, 89), ["VeriFi AI"])
fill_text(shape(s1, 90), [[("Stopping deepfake-driven payment fraud ", {}),
                           ("before the money moves", {"b": True, "color": "3C4043"})]])
place(shape(s1, 90), w=7.9)
team = shape(s1, 92)
fill_text(team, ["VeriFi AI", "FinTech & Cyber Security"])
bullet_p = shape(Presentation(SRC).slides[0], 92).text_frame._txBody.findall(qn("a:p"))[2]
for names in ["Priyanshi Chittora · Kashish Yadav", "Hardik Tailor · Jayant Kumawat"]:
    p = copy.deepcopy(bullet_p)
    rs = p.findall(qn("a:r"))
    for r in rs[1:]:
        p.remove(r)
    rs[0].find(qn("a:t")).text = names
    team.text_frame._txBody.append(p)
place(team, w=6.2)

# "built on" pill under the title
pill_t = "Built on Google Cloud · Gemini · Firebase"
pw = 0.24 + 3 * 0.36 + 0.1 + tw(pill_t, 15, True) + 0.25
box(s1, 1.04, 2.98, pw, 0.56, fill=WHITE, line=LINE, radius=0.5, lift=True, name="Built-on pill")
for i, (nm, col) in enumerate([("SiGooglecloud", BLUE), ("SiGooglegemini", RED), ("SiFirebase", YELLOW)]):
    icon(s1, nm, col, 1.28 + i * 0.36, 3.11, 0.3)
text(s1, 1.28 + 3 * 0.36 + 0.1, 2.98, pw, 0.56, [pill_t], size=15, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
# Trace → Decide → Confirm → Respond
x = 1.04
for i, (word, col, ic) in enumerate([("Trace", BLUE, "MdAccountTree"), ("Decide", RED, "MdRule"),
                                     ("Confirm", YELLOW, "MdFingerprint"), ("Respond", GREEN, "MdGavel")]):
    badge(s1, x, 3.98, 0.62, col, ic)
    text(s1, x + 0.74, 3.98, 1.5, 0.62, [word], size=18, bold=True, color=TEXT_OF[col], anchor=MSO_ANCHOR.MIDDLE)
    x += 0.74 + tw(word, 18, True) + 0.12
    if i < 3:
        icon(s1, "MdChevronRight", ICON_INK, x, 4.12, 0.34)
        x += 0.46

# =====================================================================
# 2. Problem statement
# =====================================================================
remove(s2, 101, 102, 103, 104, 105)
google_badge(s2)
text(s2, 1.08, 2.62, 17.9, 0.5, [
    [("A deepfaked CFO can talk honest staff into wiring crores, ", {}),
     ("and nothing connects the call to the payment.", {"b": True, "color": TEXT_OF[RED]})],
], size=22, color=INK)
label(s2, 1.08, 3.42, "How the fraud unfolds", w=6)
text(s2, 4.85, 3.44, 3, 0.3, ["synthetic scenario"], size=12, color=SOFT)

steps = [
    ("MdEmail", BLUE, "09:12 · GMAIL", "Lookalike email",
     "From kestrelinfra-group.co, a lookalike domain, so **SPF, DKIM, DMARC all pass.**"),
    ("MdVideocam", RED, "11:30 · VIDEO", "Deepfake “CFO”",
     "**Face and voice cloned:** “Confidential deal. Pay today, tell no one.”"),
    ("MdAccountBalance", YELLOW, "11:52 · ERP", "Brand-new account",
     "**₹5 Cr** to “Meridian Advisory LLP”, an account **never paid before.**"),
    ("MdCurrencyRupee", GREEN, "12:05 · BANK", "Money leaves",
     "The convinced AP maker releases it. **Recovery needs reporting within hours.**"),
]
CW, CG, CY, CH = 2.78, 0.35, 3.9, 2.9
for i, (ic, col, when, title, desc) in enumerate(steps):
    x = 0.94 + i * (CW + CG)
    box(s2, x, CY, CW, CH, lift=True, name=f"Step {i + 1}")
    badge(s2, x + 0.25, CY + 0.25, 0.72, col, ic)
    text(s2, x + 1.1, CY + 0.25, CW - 1.2, 0.72, [when], size=11.5, bold=True, color=TEXT_OF[col],
         anchor=MSO_ANCHOR.MIDDLE, spc=0.8)
    text(s2, x + 0.25, CY + 1.15, CW - 0.45, 0.4, [title], size=17, bold=True)
    text(s2, x + 0.25, CY + 1.6, CW - 0.45, 1.2, [desc], size=13.5, color=MUTED, line=1.08)
    if i < 3:
        arrow(s2, [(x + CW + 0.05, CY + 0.61), (x + CW + CG - 0.05, CY + 0.61)], color=SOFT, width=2)

# public-record stat card (dark, Google grey 900)
SX, SW = 13.42, 5.64
box(s2, SX, 3.42, SW, 3.38, fill=INK, lift=True, name="Stat card")
badge(s2, SX + 0.35, 3.72, 0.62, RED, "MdWarningAmber")
text(s2, SX + 1.12, 3.72, 4.2, 0.62, ["IN THE PUBLIC RECORD"], size=12, bold=True, color="BDC1C6",
     anchor=MSO_ANCHOR.MIDDLE, spc=1.2)
text(s2, SX + 0.35, 4.45, 5.0, 0.85, ["US$25.6M"], size=50, bold=True, color=WHITE)
text(s2, SX + 0.35, 5.33, 5.0, 0.6, ["lost by Arup (Hong Kong, Jan 2024) after **one video call with a deepfaked CFO**"],
     size=13.5, color="DADCE0", line=1.05, emph=WHITE)
for j, (num, lab, col) in enumerate([("15", "transfers", BLUE), ("5", "accounts", YELLOW), ("~1 wk", "to notice", GREEN)]):
    mx = SX + 0.35 + j * 1.7
    text(s2, mx, 5.98, 1.6, 0.45, [num], size=22, bold=True, color=col)
    text(s2, mx, 6.4, 1.6, 0.3, [lab], size=11.5, color="BDC1C6")
text(s2, SX, 6.9, SW, 0.3, ["Source: South China Morning Post, May 2024"], size=11, color=SOFT, align=PP_ALIGN.RIGHT)

label(s2, 1.08, 7.3, "Target users", color=MUTED, w=5)
text(s2, 7.0, 7.32, 12.06, 0.3, [[("Buyer: ", {"b": True}),
                                  ("Indian mid-to-large enterprises with central treasury / AP", {})]],
     size=13, color=MUTED, align=PP_ALIGN.RIGHT)
personas = [
    ("MdPerson", BLUE, "AP / Treasury maker", "The deceived user. Releases payments and needs **zero friction on normal days.**"),
    ("MdVerifiedUser", RED, "CFO / Authority", "The impersonated person. **Confirms or denies in seconds,** from anywhere."),
    ("MdManageSearch", GREEN, "Fraud analyst", "Investigates holds and must **report to 1930 within the golden hour.**"),
]
PW, PG = 5.84, 0.3
for i, (ic, col, title, desc) in enumerate(personas):
    x = 0.94 + i * (PW + PG)
    box(s2, x, 7.75, PW, 1.72, lift=True, name=f"Persona {title}")
    badge(s2, x + 0.3, 8.21, 0.8, col, ic)
    text(s2, x + 1.35, 8.0, PW - 1.6, 0.4, [title], size=17, bold=True)
    text(s2, x + 1.35, 8.45, PW - 1.6, 0.9, [desc], size=14.5, color=MUTED, line=1.08)

# =====================================================================
# 3. Solution
# =====================================================================
remove(s3, 116, 117, 118, 119, 121)
google_badge(s3)
text(s3, 1.08, 2.82, 17.9, 1.0, [
    [("VeriFi AI holds a risky payment until ", {}),
     ("the person the instruction claims to come from", {"b": True, "color": TEXT_OF[BLUE]}),
     (" confirms it on their own phone, with a ", {}),
     ("passkey bound to that exact amount and account.", {"b": True})],
], size=21, line=1.12)
pillars = [
    ("01", "Trace", BLUE, "MdAccountTree",
     "Gmail and Meet signals become a **graph of where every bank account came from.**",
     [("SiGmail", "Gmail API"), ("SiGooglemeet", "Meet API"), ("SiGooglecloudspanner", "Spanner Graph")]),
    ("02", "Decide", RED, "MdRule",
     "Explicit rules plus an explainable risk score. **Routine payments pass with zero friction.**",
     [("SiGooglegemini", "Gemini"), ("SiGooglecloud", "Cloud Run")]),
    ("03", "Confirm", YELLOW, "MdFingerprint",
     "The **real authority, never the requester,** approves or denies with a passkey bound to amount + account.",
     [("SiFirebase", "Firebase FCM"), ("SiFidoalliance", "Passkeys")]),
    ("04", "Respond", GREEN, "MdGavel",
     "Tamper-evident evidence and a **1930 / cybercrime.gov.in incident pack in under two minutes.**",
     [("MdKey", "Cloud KMS"), ("SiGooglecloudstorage", "Bucket Lock"), ("SiGooglegemini", "Gemini")]),
]
PW, PG, PY, PH = 4.26, 0.36, 4.05, 4.85
for i, (num, title, col, ic, desc, techs) in enumerate(pillars):
    x = 0.94 + i * (PW + PG)
    box(s3, x, PY, PW, PH, lift=True, name=f"Pillar {title}")
    badge(s3, x + 0.35, PY + 0.35, 1.0, col, ic)
    text(s3, x + PW - 1.55, PY + 0.3, 1.25, 0.7, [num], size=34, bold=True, color="E8EAED", align=PP_ALIGN.RIGHT)
    text(s3, x + 0.35, PY + 1.55, PW - 0.6, 0.55, [title], size=25, bold=True, color=TEXT_OF[col])
    text(s3, x + 0.35, PY + 2.15, PW - 0.6, 1.2, [desc], size=14.5, color=INK, line=1.1)
    text(s3, x + 0.35, PY + 3.4, PW - 0.6, 0.28, ["POWERED BY"], size=10.5, bold=True, color=SOFT, spc=1.2)
    cx, cy = x + 0.35, PY + 3.72
    for nm, t in techs:
        w = 0.2 + 0.4 * 0.6 + 0.1 + tw(t, 12, True) + 0.2
        if cx + w > x + PW - 0.25:
            cx, cy = x + 0.35, cy + 0.52
        chip(s3, cx, cy, nm, col, t, h=0.4, size=12, fill=TINT[col], line=None)
        cx += w + 0.1
    if i < 3:  # connector bubble between pillars
        bx = x + PW + PG / 2 - 0.25
        box(s3, bx, PY + 0.6, 0.5, 0.5, fill=WHITE, line=LINE, geom=MSO_SHAPE.OVAL, lift=True)
        icon(s3, "MdArrowForward", ICON_INK, bx + 0.11, PY + 0.71, 0.28)
box(s3, 0.94, 9.2, 18.12, 0.72, fill=TINT[BLUE], radius=0.5, name="Tagline")
icon(s3, "SiGooglegemini", BLUE, 1.25, 9.38, 0.36)
text(s3, 1.8, 9.2, 17.0, 0.72, [
    [("Detectors tell you a call might be fake. ", {}),
     ("VeriFi makes sure money can't move because of it.", {"b": True, "color": TEXT_OF[BLUE]})],
], size=17, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 4. Technical approach: Google Cloud architecture
# =====================================================================
remove(s4, 132)
google_badge(s4)
fill_text(shape(s4, 131), ["ARCHITECTURE DIAGRAM / WORKFLOW"])
place(shape(s4, 131), x=1.08, y=2.62, w=9.0, h=0.4)
for p in shape(s4, 131).text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(16)
DASH, DOT = MSO_LINE_DASH_STYLE.DASH, MSO_LINE_DASH_STYLE.ROUND_DOT
# legend
lx = 10.9
for t, col, dash in [("slow loop · async", BLUE, DASH), ("fast loop · ≤ 300 ms", RED, None), ("evidence", GREEN, DOT)]:
    arrow(s4, [(lx, 2.82), (lx + 0.55, 2.82)], color=col, width=2.25, dash=dash)
    text(s4, lx + 0.65, 2.66, 2.6, 0.32, [t], size=12.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    lx += 0.65 + tw(t, 12.5) + 0.45

NW, NG, NX0 = 2.55, 0.33, 4.12        # node width, gap, first column x
COLX = [NX0 + k * (NW + NG) for k in range(4)]


def node(slide, x, y, w, h, ic, col, title, sub, d=0.62, tsize=13.5, ssize=11):
    box(slide, x, y, w, h, lift=True, name=f"Node {title}")
    badge(slide, x + 0.16, y + (h - d) / 2, d, col, ic)
    text(slide, x + d + 0.3, y + 0.08, w - d - 0.4, h - 0.16,
         [[(title, {"b": True, "sz": tsize})], [(sub, {"sz": ssize, "color": MUTED})]],
         color=INK, line=1.0, anchor=MSO_ANCHOR.MIDDLE, space_after=2)


# Google Cloud boundary
box(s4, 3.82, 3.22, 11.78, 6.83, fill="FBFCFE", line="AECBFA", dash=True, line_w=1.5, radius=0.03, name="GCP boundary")
icon(s4, "SiGooglecloud", BLUE, 4.05, 3.36, 0.36)
text(s4, 4.52, 3.33, 10.9, 0.42, [[("Google Cloud", {"b": True}),
                               ("  ·  asia-south1 (Mumbai) · ", {"color": MUTED, "sz": 12.5}),
                               ("Gemini reads, rules decide, the real CFO confirms", {"b": True, "color": TEXT_OF[BLUE], "sz": 12.5})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
# lanes
LANES = [("Slow loop · Trace", BLUE, 3.88, 2.15), ("Fast loop · Decide  (no LLM on this path)", RED, 6.17, 2.15),
         ("Trust & evidence", GREEN, 8.46, 1.45)]
for t, col, y, h in LANES:
    box(s4, 3.98, y, 11.46, h, fill=TINT[col], radius=0.06, name=f"Lane {t}")
    text(s4, 4.12, y + 0.08, 8, 0.3, [t.upper()], size=11, bold=True, color=TEXT_OF[col], spc=1.2)
A_Y, B_Y, NH = 4.27, 6.56, 1.62
lane_a = [("SiGooglepubsub", BLUE, "Pub/Sub", "Gmail push + Meet events"),
          ("MdShield", RED, "Sensitive Data Protection", "Tokenize A/c, IFSC; drop Aadhaar"),
          ("SiGooglegemini", YELLOW, "Gemini on Vertex AI", "Intent, urgency, authority claims + Web Risk"),
          ("SiGooglecloudspanner", GREEN, "Spanner Graph", "Instruction Lineage Graph")]
lane_b = [("SiGooglecloud", BLUE, "Cloud Run API", "FastAPI · /payments/evaluate"),
          ("MdRule", RED, "Policy engine", "Rules R1–R7 + bounded risk score"),
          ("MdTraffic", YELLOW, "Assurance tier", "T0 pass · T1 confirm · T2 two + wait"),
          ("SiFirebase", GREEN, "Firebase Cloud Messaging", "Push to the real authority")]
for k, args in enumerate(lane_a):
    node(s4, COLX[k], A_Y, NW, NH, *args)
for k, args in enumerate(lane_b):
    node(s4, COLX[k], B_Y, NW, NH, *args)
for k in range(3):
    arrow(s4, [(COLX[k] + NW + 0.03, A_Y + NH / 2), (COLX[k + 1] - 0.03, A_Y + NH / 2)], color=BLUE, width=2, dash=DASH)
    arrow(s4, [(COLX[k] + NW + 0.03, B_Y + NH / 2), (COLX[k + 1] - 0.03, B_Y + NH / 2)], color=RED, width=2)
# provenance facts: Spanner Graph → Policy engine
sx, px = COLX[3] + NW / 2, COLX[1] + NW / 2
arrow(s4, [(sx, A_Y + NH), (sx, 6.4), (px, 6.4), (px, B_Y)], color=GREEN, width=2)
text(s4, COLX[2] - 0.2, 6.13, 2.6, 0.26, ["provenance facts · 150 ms"], size=10.5, bold=True, color=TEXT_OF[GREEN],
     align=PP_ALIGN.CENTER)
# trust & evidence lane (compact nodes)
C_Y, CNH = 8.8, 1.0
lane_c = [("MdStorage", BLUE, "Cloud SQL", "System of record"), ("MdKey", RED, "Cloud KMS", "Signs the evidence chain"),
          ("SiGooglebigquery", YELLOW, "BigQuery", "Fraud analytics"),
          ("SiGooglecloudstorage", GREEN, "Bucket Lock", "WORM evidence store")]
for k, (ic, col, t, sub) in enumerate(lane_c):
    node(s4, COLX[k], C_Y, NW, CNH, ic, col, t, sub, d=0.52, tsize=13, ssize=10.5)
arrow(s4, [(px, B_Y + NH), (px, C_Y)], color=GREEN, width=2, dash=DOT)
for k in range(3):
    arrow(s4, [(COLX[k] + NW + 0.03, C_Y + CNH / 2), (COLX[k + 1] - 0.03, C_Y + CNH / 2)], color=GREEN,
          width=1.75, dash=DOT)

# sources (left) and people (right)
label(s4, 0.94, 3.3, "Sources", color=MUTED, size=12, w=2.6)
label(s4, 15.98, 3.3, "People", color=MUTED, size=12, w=3)
SRCX, SRCW = 0.94, 2.55
src = [("SiGmail", RED, "Gmail", "Finance mailboxes", 4.12),
       ("SiGooglemeet", GREEN, "Google Meet", "Participants, transcripts", 5.16),
       ("MdAccountBalance", BLUE, "ERP / Treasury", "Payment release hook", 6.86)]
for ic, col, t, sub, y in src:
    node(s4, SRCX, y, SRCW, 0.92, ic, col, t, sub, d=0.5, tsize=13, ssize=10.5)
arrow(s4, [(SRCX + SRCW + 0.03, 4.58), (COLX[0] - 0.03, 4.58)], color=BLUE, width=2, dash=DASH)
arrow(s4, [(SRCX + SRCW + 0.03, 5.62), (COLX[0] - 0.03, 5.62)], color=BLUE, width=2, dash=DASH)
arrow(s4, [(SRCX + SRCW + 0.03, 7.32), (COLX[0] - 0.03, 7.32)], color=RED, width=2)
text(s4, 0.94, 7.9, 2.55, 0.7, ["**release or hold** returned to the ERP, HMAC-signed"], size=10.5, color=SOFT)
PPX, PPW = 15.98, 3.08
people = [("MdDashboard", BLUE, "Analyst console", "Gemini copilot explains holds", A_Y, NH),
          ("MdPhoneIphone", YELLOW, "CFO's phone", "Passkey: Approve / Deny", B_Y, NH),
          ("MdDescription", RED, "1930 incident pack", "cybercrime.gov.in + bank recall", C_Y, CNH)]
for ic, col, t, sub, y, h in people:
    big = h > 1.2
    node(s4, PPX, y, PPW, h, ic, col, t, sub, d=0.62 if big else 0.52, tsize=13.5 if big else 13,
         ssize=11 if big else 10.5)
arrow(s4, [(COLX[3] + NW + 0.03, A_Y + NH / 2), (PPX - 0.03, A_Y + NH / 2)], color=BLUE, width=2, dash=DASH)
arrow(s4, [(COLX[3] + NW + 0.03, B_Y + NH / 2), (PPX - 0.03, B_Y + NH / 2)], color=RED, width=2)
arrow(s4, [(COLX[3] + NW + 0.03, C_Y + CNH / 2), (PPX - 0.03, C_Y + CNH / 2)], color=GREEN, width=1.75, dash=DOT)

# =====================================================================
# 5. Tech stack
# =====================================================================
remove(s5, *range(144, 160))
google_badge(s5)
cols = [
    ("Frontend", BLUE, "MdWeb", [("SiNextdotjs", "Next.js + TypeScript"), ("MdPhoneIphone", "Approver PWA"),
                                 ("MdDashboard", "Analyst Console")]),
    ("Backend", RED, "MdDns", [("SiFastapi", "FastAPI · Python"), ("SiGooglepubsub", "Pub/Sub workers"),
                               ("SiGooglegemini", "Google Gen AI SDK")]),
    ("Database", YELLOW, "MdStorage", [("SiGooglecloudspanner", "Spanner Graph"), ("MdStorage", "Cloud SQL Postgres"),
                                       ("SiGooglecloudstorage", "Cloud Storage")]),
    ("Auth & API", GREEN, "MdFingerprint", [("SiWebauthn", "Passkeys (WebAuthn)"), ("SiFirebase", "Firebase Auth + FCM"),
                                            ("MdKey", "Cloud KMS signing")]),
    ("Deployment", BLUE, "MdCloudUpload", [("SiGooglecloud", "Cloud Run · Mumbai"), ("SiFirebase", "Firebase App Hosting"),
                                           ("MdBuild", "Cloud Build CI/CD")]),
]
KW, KG, KY, KH = 3.384, 0.3, 2.85, 3.55
for i, (t, col, ic, items) in enumerate(cols):
    x = 0.94 + i * (KW + KG)
    box(s5, x, KY, KW, KH, lift=True, name=f"Stack {t}")
    badge(s5, x + 0.3, KY + 0.3, 0.78, col, ic)
    text(s5, x + 1.25, KY + 0.3, KW - 1.4, 0.78, [t.upper()], size=15, bold=True, color=TEXT_OF[col],
         anchor=MSO_ANCHOR.MIDDLE, spc=1)
    for j, (inm, it) in enumerate(items):
        iy = KY + 1.38 + j * 0.68
        box(s5, x + 0.3, iy, KW - 0.6, 0.54, fill=GREY_BG, radius=0.3, name=f"Item {it}")
        icon(s5, inm, ICON_INK, x + 0.45, iy + 0.12, 0.3)
        text(s5, x + 0.9, iy, KW - 1.25, 0.54, [it], size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)

RY, RH = 6.7, 3.25
# Google AI layer
box(s5, 0.94, RY, 8.91, RH, lift=True, name="AI layer card")
badge(s5, 1.24, RY + 0.3, 0.78, BLUE, "SiGooglegemini")
text(s5, 2.2, RY + 0.28, 7.4, 0.45, ["GOOGLE AI LAYER"], size=16, bold=True, color=TEXT_OF[BLUE], spc=1)
text(s5, 2.2, RY + 0.7, 7.4, 0.4, ["Gemini does the reading; **explicit rules make the decision.**"], size=13.5,
     color=MUTED)
ai = [("SiGooglegemini", BLUE, "Gemini on Vertex AI", "structured extraction of intent, urgency, authority claims"),
      ("MdPsychology", RED, "Gemini analyst copilot", "explains each hold; drafts the 1930 pack + bank recall"),
      ("MdShield", GREEN, "Sensitive Data Protection + Web Risk", "tokenize PII; flag lookalike domains")]
for j, (ic, col, k, v) in enumerate(ai):
    ay = RY + 1.3 + j * 0.62
    badge(s5, 1.24, ay + 0.04, 0.46, col, ic)
    text(s5, 1.9, ay, 7.8, 0.56, [[(k, {"b": True}), ("  " + v, {"color": MUTED, "sz": 13})]], size=14.5,
         anchor=MSO_ANCHOR.MIDDLE)
# configurable by design
CX0 = 10.15
box(s5, CX0, RY, 8.91, RH, lift=True, name="Config card")
badge(s5, CX0 + 0.3, RY + 0.3, 0.78, GREEN, "MdTune")
text(s5, CX0 + 1.26, RY + 0.28, 7.4, 0.45, ["CONFIGURABLE BY DESIGN"], size=16, bold=True, color=TEXT_OF[GREEN], spc=1)
text(s5, CX0 + 1.26, RY + 0.7, 7.4, 0.4, ["**Google by default.** Every layer sits behind an adapter, **swapped in config.**"],
     size=13.5, color=MUTED)
cfg = [("LLM", "SiGooglegemini", "Gemini · Vertex AI", "Claude · local model"),
       ("Graph", "SiGooglecloudspanner", "Spanner Graph", "Neo4j AuraDB"),
       ("Evidence", "MdKey", "Cloud KMS + Lock", "AWS KMS + S3 Lock"),
       ("Push", "SiFirebase", "Firebase FCM", "Web Push")]
for j, (layer, ic, g, alt) in enumerate(cfg):
    ry = RY + 1.27 + j * 0.47
    text(s5, CX0 + 0.35, ry, 1.3, 0.4, [layer], size=13.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s5, CX0 + 1.6, ry + 0.02, 3.25, 0.36, fill=TINT[BLUE], radius=0.5)
    icon(s5, ic, BLUE, CX0 + 1.74, ry + 0.07, 0.26)
    text(s5, CX0 + 2.1, ry + 0.02, 2.7, 0.36, [g], size=12.5, bold=True, color=TEXT_OF[BLUE], anchor=MSO_ANCHOR.MIDDLE)
    icon(s5, "MdSwapHoriz", ICON_INK, CX0 + 4.98, ry + 0.04, 0.32)
    box(s5, CX0 + 5.42, ry + 0.02, 3.2, 0.36, fill=GREY_BG, radius=0.5)
    text(s5, CX0 + 5.6, ry + 0.02, 3.0, 0.36, [alt], size=12.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 6. Impact & scalability
# =====================================================================
remove(s6, *range(169, 184))
google_badge(s6)
left = [("MdShield", RED, "Real impact", "Stops crore-scale deepfake and BEC fraud **before the money leaves,** not days later."),
        ("MdGroups", BLUE, "More users", "Routine payments pass untouched; **only risky ones need one tap** from the real CFO."),
        ("MdRocketLaunch", GREEN, "Built to scale", "Serverless on Cloud Run: **≤ 300 ms decisions,** no LLM on the fast path.")]
for i, (ic, col, t, d) in enumerate(left):
    y = 3.0 + i * 2.3
    badge(s6, 1.0, y, 0.95, col, ic)
    text(s6, 2.25, y - 0.02, 5.2, 0.5, [t], size=23, bold=True)
    text(s6, 2.25, y + 0.52, 5.0, 1.3, [d], size=16.5, color=MUTED, line=1.1)

PX0, PWD = 7.9, 11.16
box(s6, PX0, 2.85, PWD, 7.1, lift=True, name="Rollout card")
label(s6, PX0 + 0.45, 3.12, "Rollout path", w=5)
stops = [("MdLocationOn", BLUE, "UDAIPUR", "**Shadow-mode pilots** with local enterprise finance teams."),
         ("MdMap", RED, "RAJASTHAN", "**MSME exporters, cooperative banks** and state enterprises."),
         ("MdPublic", GREEN, "INDIA", "**Every NEFT / RTGS / IMPS release,** wired to 1930 reporting.")]
centers = [PX0 + 1.95 + k * 3.63 for k in range(3)]
for a, b in zip(centers, centers[1:]):
    arrow(s6, [(a + 0.7, 4.17), (b - 0.7, 4.17)], color=LINE, width=2.5, dash=DASH)
for (ic, col, t, d), c in zip(stops, centers):
    badge(s6, c - 0.55, 3.62, 1.1, col, ic)
    text(s6, c - 1.6, 4.9, 3.2, 0.5, [t], size=25, bold=True, color=TEXT_OF[col], align=PP_ALIGN.CENTER)
    text(s6, c - 1.55, 5.45, 3.1, 1.0, [d], size=15, color=MUTED, align=PP_ALIGN.CENTER, line=1.08)
label(s6, PX0 + 0.45, 6.95, "Pilot plan", w=5)
plan = [("MdVisibility", BLUE, "60-day shadow mode", "Runs on the customer's own data; **never blocks a payment.**"),
        ("MdTune", YELLOW, "Calibrate", "Tune weights and thresholds **per customer, from real outcomes.**"),
        ("MdTaskAlt", GREEN, "Go live", "**Gate the payment-release step** through an ERP API hook.")]
QW, QG = 3.3, 0.28
for k, (ic, col, t, d) in enumerate(plan):
    x = PX0 + 0.45 + k * (QW + QG)
    box(s6, x, 7.42, QW, 2.05, fill=TINT[col], radius=0.08, name=f"Plan {t}")
    badge(s6, x + 0.25, 7.65, 0.6, col, ic)
    text(s6, x + 1.0, 7.65, QW - 1.1, 0.6, [t], size=15.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s6, x + 0.25, 8.45, QW - 0.45, 0.9, [d], size=13.5, color=INK, line=1.08)
    if k < 2:
        icon(s6, "MdChevronRight", ICON_INK, x + QW - 0.02, 8.29, 0.32)

prs.save(OUT)
print("wrote", OUT)
